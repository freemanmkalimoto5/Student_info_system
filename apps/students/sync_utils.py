"""
Bulk helpers used by the Import page:

  export_students_xlsx()   -> every student in ONE universal Excel file
  replace_from_table(...)  -> take an edited copy of that file, compare it
                              with what is in the system, then
                                * edit only the students that really changed
                                * silently skip students that are identical
                                * add students that are not in the system yet

No "possible duplicate" notification is created here on purpose: matching an
existing student is the whole point of this feature.
"""
import csv
import io
import uuid
from datetime import datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction

from apps.accounts.services import log_action, sync_all_parent_accounts
from apps.pocketmoney.services import add_transaction

from .models import Student
from .services import renumber_students, sort_students

# Column order of the exported file (also what the replace step understands).
EXPORT_COLUMNS = [
    'student_number',
    'first_name', 'middle_name', 'last_name', 'date_of_birth',
    'grade_class', 'status',
    'father_name', 'father_phone', 'father_whatsapp',
    'mother_name', 'mother_phone', 'mother_whatsapp',
    'guardian_name', 'guardian_phone', 'guardian_relationship',
    'parish', 'full_address', 'region', 'district', 'ward',
]

# Optional text columns -- a blank cell means "empty".
TEXT_FIELDS = [
    'middle_name',
    'father_name', 'father_phone', 'father_whatsapp',
    'mother_name', 'mother_phone', 'mother_whatsapp',
    'guardian_name', 'guardian_phone', 'guardian_relationship',
    'parish', 'full_address', 'region', 'district', 'ward',
]
PHONE_FIELDS = {
    'father_phone', 'father_whatsapp', 'mother_phone', 'mother_whatsapp', 'guardian_phone',
}
# Changing any of these means the parent logins must be re-synced.
PARENT_FIELDS = {
    'father_name', 'father_phone', 'father_whatsapp',
    'mother_name', 'mother_phone', 'mother_whatsapp',
    'guardian_name', 'guardian_phone', 'guardian_relationship',
}
# Changing any of these means the QR code (which carries this data) is out of date.
QR_FIELDS = {
    'first_name', 'middle_name', 'last_name', 'grade_class', 'status', 'parish',
}

CLASS_MAP = {key: key for key, _ in Student.CLASS_CHOICES}
CLASS_MAP.update({label.lower(): key for key, label in Student.CLASS_CHOICES})
STATUS_MAP = {key: key for key, _ in Student.STATUS_CHOICES}
STATUS_MAP.update({label.lower(): key for key, label in Student.STATUS_CHOICES})


class RowError(Exception):
    """A single row that cannot be used; the message is shown to the admin."""


# ======================================================================
# EXPORT
# ======================================================================
def export_students_xlsx():
    """Returns the bytes of an .xlsx file holding every student."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        raise ValueError(
            "This server can't create Excel files yet -- the 'openpyxl' package "
            "isn't installed. Run: pip install openpyxl"
        )

    wb = Workbook()
    ws = wb.active
    ws.title = 'Students'
    ws.append(EXPORT_COLUMNS)

    head_font = Font(name='Arial', bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', start_color='1F3B5E')
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(vertical='center')

    for student in sort_students(Student.objects.all()):
        ws.append([getattr(student, column) for column in EXPORT_COLUMNS])

    body_font = Font(name='Arial')
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            name = EXPORT_COLUMNS[cell.column - 1]
            if name in PHONE_FIELDS or name == 'student_number':
                cell.number_format = '@'          # keeps leading zeros in phone numbers
            elif name == 'date_of_birth':
                cell.number_format = 'yyyy-mm-dd'

    # Comfortable column widths.
    for index, name in enumerate(EXPORT_COLUMNS, start=1):
        longest = max(
            (len(str(c.value)) for c in ws[get_column_letter(index)] if c.value is not None),
            default=10,
        )
        ws.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), 40)
    ws.freeze_panes = 'A2'

    # Drop-downs so edits to class / status stay valid.
    last_row = max(ws.max_row, 2) + 500
    for column, options in (
        ('grade_class', ','.join(key for key, _ in Student.CLASS_CHOICES)),
        ('status', ','.join(key for key, _ in Student.STATUS_CHOICES)),
    ):
        letter = get_column_letter(EXPORT_COLUMNS.index(column) + 1)
        validation = DataValidation(type='list', formula1=f'"{options}"', allow_blank=True)
        ws.add_data_validation(validation)
        validation.add(f'{letter}2:{letter}{last_row}')

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ======================================================================
# REPLACE EXISTING DETAILS
# ======================================================================
def _clean(value):
    if value is None:
        return ''
    return str(value).replace('\r\n', '\n').strip()


def _fold(value):
    """Lower-case + single spaces, so 'John  DOE' matches 'john doe'."""
    return ' '.join((value or '').split()).lower()


def _same_phone(old, new):
    old, new = _clean(old), _clean(new)
    # Excel often drops the leading zero of a number typed as a plain number.
    return old == new or old == '0' + new


def _parse_date(text):
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _choice(text, mapping, label):
    key = _fold(text).replace('_', ' ')
    if key in mapping:
        return mapping[key]
    allowed = ', '.join(sorted(set(mapping.values())))
    raise RowError(f"{label} '{text}' is not valid (use one of: {allowed}).")


def _money(text):
    text = _clean(text).replace(',', '')
    if not text:
        return None
    try:
        amount = Decimal(text)
    except InvalidOperation:
        raise RowError("initial_pocket_money must be a number.")
    return amount if amount > 0 else None


def _read_rows(csv_text):
    rows = list(csv.reader(io.StringIO(csv_text)))
    if not rows:
        raise ValueError("The file appears to be empty.")

    header = [h.strip().lower().replace(' ', '_') for h in rows[0]]
    missing = [c for c in ('first_name', 'last_name', 'date_of_birth') if c not in header]
    if missing:
        raise ValueError(
            "The file is missing the column(s): " + ", ".join(missing) +
            ". Use the file from 'Download a file' as your starting point."
        )

    data = []
    for line_no, raw in enumerate(rows[1:], start=2):
        if not any(cell.strip() for cell in raw):
            continue
        raw = list(raw) + [''] * (len(header) - len(raw))      # short rows = blank cells
        data.append((line_no, {header[i]: raw[i].strip() for i in range(len(header)) if header[i]}))
    return header, data


def _find_match(row, values, by_number, by_identity):
    first_key, last_key = _fold(values['first_name']), _fold(values['last_name'])

    # 1) Same student number, as long as it still looks like the same person.
    number = _clean(row.get('student_number')).upper()
    if number:
        student = by_number.get(number)
        if student and (
            student.date_of_birth == values['date_of_birth']
            or (_fold(student.first_name) == first_key and _fold(student.last_name) == last_key)
        ):
            return student

    # 2) Same first name + last name + date of birth.
    candidates = by_identity.get((first_key, last_key, values['date_of_birth']), [])
    if candidates:
        middle = _fold(values.get('middle_name'))
        for student in candidates:
            if _fold(student.middle_name) == middle:
                return student
        return candidates[0]
    return None


def _update(student, values, user):
    changed = []
    for field, new in values.items():
        old = getattr(student, field)
        if field in PHONE_FIELDS:
            same = _same_phone(old, new)
        elif field in ('date_of_birth', 'grade_class', 'status'):
            same = old == new
        else:
            same = _clean(old) == _clean(new)
        if not same:
            setattr(student, field, new)
            changed.append(field)

    if not changed:
        return 'unchanged', student, []

    if 'date_of_birth' in changed:
        student.age = Student.calculate_age(student.date_of_birth)
    student.save()
    log_action(student, 'updated', user)
    if PARENT_FIELDS.intersection(changed):
        sync_all_parent_accounts(student, notify=False)
    return 'updated', student, changed


def _create(values, row, user):
    if 'grade_class' not in values:
        raise RowError("grade_class is required to add a new student.")

    student = Student(
        student_number=f"TMP-{uuid.uuid4().hex[:10]}",   # real number is set by renumber_students()
        added_by=user,
        **values,
    )
    student.save()
    sync_all_parent_accounts(student, notify=False)
    log_action(student, 'created', user)

    amount = _money(row.get('initial_pocket_money'))
    if amount:
        add_transaction(
            student, 'deposit', amount,
            'Initial pocket money at registration', user, notify=False,
        )
    return 'created', student, []


def _process_row(row, columns, user, by_number, by_identity, used, created_keys):
    first, last = _clean(row.get('first_name')), _clean(row.get('last_name'))
    if not first or not last:
        raise RowError("first_name and last_name are required.")
    dob = _parse_date(_clean(row.get('date_of_birth')))
    if dob is None:
        raise RowError("date_of_birth is missing or not a valid date (use YYYY-MM-DD).")

    values = {'first_name': first, 'last_name': last, 'date_of_birth': dob}
    if 'grade_class' in columns and _clean(row.get('grade_class')):
        values['grade_class'] = _choice(_clean(row['grade_class']), CLASS_MAP, 'grade_class')
    if 'status' in columns and _clean(row.get('status')):
        values['status'] = _choice(_clean(row['status']), STATUS_MAP, 'status')
    for field in TEXT_FIELDS:
        if field in columns:
            values[field] = _clean(row.get(field)) or None

    student = _find_match(row, values, by_number, by_identity)
    if student is not None:
        if student.pk in used:
            raise RowError("This student appears more than once in the file.")
        return _update(student, values, user)

    if (_fold(first), _fold(last), dob) in created_keys:
        raise RowError("This student appears more than once in the file.")
    return _create(values, row, user)


def replace_from_table(csv_text, user):
    """
    csv_text: the uploaded file already converted to CSV text
              (see import_utils.read_uploaded_table_as_csv_text).
    Returns {'created': [...], 'updated': [...], 'unchanged': int, 'skipped': [...]}.
    Raises ValueError (message safe to show) if the file is unusable as a whole.
    """
    header, rows = _read_rows(csv_text)
    columns = set(header)

    students = list(Student.objects.all())
    by_number = {s.student_number.upper(): s for s in students}
    by_identity = {}
    for s in students:
        by_identity.setdefault((_fold(s.first_name), _fold(s.last_name), s.date_of_birth), []).append(s)

    used, created_keys = set(), set()
    result = {'created': [], 'updated': [], 'unchanged': 0, 'skipped': []}
    qr_pks = []

    for line_no, row in rows:
        try:
            with transaction.atomic():           # one row failing never affects the others
                kind, student, changed = _process_row(
                    row, columns, user, by_number, by_identity, used, created_keys
                )
        except RowError as exc:
            result['skipped'].append({'row': line_no, 'reason': str(exc)})
            continue
        except Exception as exc:
            result['skipped'].append({'row': line_no, 'reason': f"Could not be saved ({exc})."})
            continue

        used.add(student.pk)
        if kind == 'created':
            created_keys.add((_fold(student.first_name), _fold(student.last_name), student.date_of_birth))
            result['created'].append({'pk': student.pk, 'name': student.full_name})
            qr_pks.append(student.pk)
        elif kind == 'updated':
            result['updated'].append({
                'pk': student.pk, 'name': student.full_name,
                'changed': [f.replace('_', ' ') for f in changed],
            })
            if QR_FIELDS.intersection(changed):
                qr_pks.append(student.pk)
        else:
            result['unchanged'] += 1

    touched = [e['pk'] for e in result['created']] + [e['pk'] for e in result['updated']]
    if touched:
        renumber_students()        # class / name changes and new students need the numbering redone
        numbers = dict(Student.objects.filter(pk__in=touched).values_list('pk', 'student_number'))
        for entry in result['created'] + result['updated']:
            entry['number'] = numbers.get(entry['pk'], '')
        _refresh_qr_codes(qr_pks)

    return result


def _refresh_qr_codes(pks):
    """The QR carries the student's data, so rebuild it where that data changed."""
    if not pks:
        return
    try:
        from apps.qrcodes.services import save_card_qr
    except ImportError:
        return
    for student in Student.objects.filter(pk__in=pks):
        try:
            save_card_qr(student)
        except Exception:
            continue      # a QR problem must never undo the data update (System Check will flag it)
