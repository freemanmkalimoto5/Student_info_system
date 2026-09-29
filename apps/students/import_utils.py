"""
Reads an uploaded file (.csv, .xlsx, .xlsm, or the older .xls) and turns it
into plain CSV text, so the rest of the import code (run_import_job /
_parse_and_validate_rows in services.py) never needs to change -- it always
just gets CSV text, no matter what format the user actually uploaded.
"""
import csv
import datetime
import io

SUPPORTED_EXTENSIONS = ('.csv', '.xlsx', '.xlsm', '.xls')


def _cell_to_text(value):
    """Excel cells can hold real dates/numbers -- turn them into the same
    plain text the CSV importer already expects."""
    if value is None:
        return ''
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows_to_csv_text(rows):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    for row in rows:
        writer.writerow(row)
    return buffer.getvalue()


def _read_xlsx(file_bytes):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError(
            "This server can't read .xlsx files yet -- the 'openpyxl' "
            "package isn't installed. Run: pip install openpyxl "
            "-- or save the file as .csv instead."
        )
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = wb.active
    rows = [[_cell_to_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    wb.close()
    return rows


def _read_xls(file_bytes):
    try:
        import xlrd
    except ImportError:
        raise ValueError(
            "This server can't read the old .xls format -- the 'xlrd' "
            "package isn't installed. Run: pip install xlrd -- or "
            "open the file in Excel and save it as .xlsx or .csv instead."
        )
    book = xlrd.open_workbook(file_contents=file_bytes)
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                dt = xlrd.xldate_as_datetime(cell.value, book.datemode)
                row.append(dt.date().isoformat())
            else:
                row.append(_cell_to_text(cell.value))
        rows.append(row)
    return rows


def read_uploaded_table_as_csv_text(uploaded_file, filename):
    """
    uploaded_file: Django's UploadedFile (from request.FILES).
    filename: the original filename, used only to tell the format apart.
    Returns CSV text ready for the existing import pipeline.
    Raises ValueError with a message safe to show the user.
    """
    name = (filename or '').lower()
    file_bytes = uploaded_file.read()

    if name.endswith('.csv'):
        return file_bytes.decode('utf-8-sig')

    if name.endswith('.xlsx') or name.endswith('.xlsm'):
        rows = _read_xlsx(file_bytes)
    elif name.endswith('.xls'):
        rows = _read_xls(file_bytes)
    else:
        raise ValueError(
            "Unsupported file type. Please upload a .csv, .xlsx, .xlsm or .xls file."
        )

    rows = [r for r in rows if any(cell.strip() for cell in r)]   # drop fully-empty rows
    if not rows:
        raise ValueError("The file appears to be empty.")
    return _rows_to_csv_text(rows)
